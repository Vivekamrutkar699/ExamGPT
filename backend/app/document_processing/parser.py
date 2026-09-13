import fitz  # PyMuPDF
import pdfplumber
from docx import Document as DocxDocument
from pptx import Presentation as PptxPresentation
from PIL import Image


class DocumentParser:
    """
    Core parsing engine supporting multiple formats:
    PDF, DOCX, PPTX, and images.
    """

    def parse_pdf(self, file_path: str) -> list[dict]:
        """
        Extract text from PDF pages using PyMuPDF.

        pdfplumber is additionally used to extract tables.
        """

        pages_content = []

        # ---------------------------------------------------------
        # 1. Extract text using PyMuPDF
        # ---------------------------------------------------------
        with fitz.open(file_path) as doc:
            for page_idx, page in enumerate(doc):

                text = page.get_text()

                # Handle scanned/empty pages gracefully
                if not text.strip():
                    text = "[Scanned Page / Unparsed Content]"

                pages_content.append(
                    {
                        "page": page_idx + 1,
                        "text": text,
                        "tables": [],
                    }
                )

        # ---------------------------------------------------------
        # 2. Extract tables using pdfplumber
        # ---------------------------------------------------------
        try:
            with pdfplumber.open(file_path) as pdf:

                for idx, page in enumerate(pdf.pages):

                    if idx >= len(pages_content):
                        continue

                    tables = page.extract_tables()

                    if tables:
                        pages_content[idx]["tables"] = tables

        except Exception as e:
            # Table extraction failure should not stop PDF processing
            print(
                f"Warning during PDF table extraction: {e}"
            )

        return pages_content

    def parse_docx(self, file_path: str) -> list[dict]:
        """
        Extract structural paragraphs from DOCX files.

        DOCX does not expose reliable page boundaries, so paragraphs
        are grouped into approximate pages.
        """

        doc = DocxDocument(file_path)

        paragraphs = []

        for para in doc.paragraphs:

            if para.text.strip():
                paragraphs.append(para.text)

        pages_content = []

        temp_text = []
        page_num = 1

        for idx, text in enumerate(paragraphs):

            temp_text.append(text)

            # Approximate page boundary every 10 paragraphs
            if (idx + 1) % 10 == 0 or (idx + 1) == len(paragraphs):

                pages_content.append(
                    {
                        "page": page_num,
                        "text": "\n".join(temp_text),
                        "tables": [],
                    }
                )

                temp_text = []
                page_num += 1

        return pages_content

    def parse_pptx(self, file_path: str) -> list[dict]:
        """
        Extract text from PowerPoint slides.
        Each slide is treated as one page.
        """

        prs = PptxPresentation(file_path)

        pages_content = []

        for slide_idx, slide in enumerate(prs.slides):

            slide_text = []

            for shape in slide.shapes:

                if hasattr(shape, "text_frame") and shape.text_frame:

                    for paragraph in shape.text_frame.paragraphs:

                        if paragraph.text.strip():
                            slide_text.append(paragraph.text)

            pages_content.append(
                {
                    "page": slide_idx + 1,
                    "text": "\n".join(slide_text),
                    "tables": [],
                }
            )

        return pages_content

    def parse_image(self, file_path: str) -> list[dict]:
        """
        Extract text from image files using Tesseract OCR.

        If pytesseract or Tesseract is unavailable, a placeholder
        is returned instead of crashing the entire pipeline.
        """

        try:

            import pytesseract

            img = Image.open(file_path)

            text = pytesseract.image_to_string(img)

            if not text.strip():
                text = "[Image contains no extractable text]"

        except ImportError:

            text = (
                "[OCR library pytesseract is not installed. "
                "Local image parsing is disabled.]"
            )

        except Exception as e:

            text = f"[OCR image parsing failed: {e}]"

        return [
            {
                "page": 1,
                "text": text,
                "tables": [],
            }
        ]

    def extract_text(
        self,
        file_path: str,
        file_type: str,
    ) -> list[dict]:
        """
        Route a file to the appropriate parser based on its extension.
        """

        ftype = file_type.lower().strip(".")

        # ---------------------------------------------------------
        # PDF
        # ---------------------------------------------------------
        if ftype == "pdf":
            return self.parse_pdf(file_path)

        # ---------------------------------------------------------
        # DOCX / DOC
        # ---------------------------------------------------------
        elif ftype in {"docx", "doc"}:
            return self.parse_docx(file_path)

        # ---------------------------------------------------------
        # PPTX / PPT
        # ---------------------------------------------------------
        elif ftype in {"pptx", "ppt"}:
            return self.parse_pptx(file_path)

        # ---------------------------------------------------------
        # Images
        # ---------------------------------------------------------
        elif ftype in {"png", "jpg", "jpeg"}:
            return self.parse_image(file_path)

        # ---------------------------------------------------------
        # ZIP
        # ---------------------------------------------------------
        elif ftype == "zip":
            raise ValueError(
                "ZIP files must be extracted before parsing single files."
            )

        # ---------------------------------------------------------
        # Unsupported format
        # ---------------------------------------------------------
        else:
            raise ValueError(
                f"No parsing handler registered for format: .{ftype}"
            )


# Singleton instance
document_parser = DocumentParser()